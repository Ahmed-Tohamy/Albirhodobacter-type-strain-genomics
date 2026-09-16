#!/usr/bin/env python3
"""Append Table S48 (plasmid versus whole-genome comparison) to Additional file 1.
Input: the S47 workbook; output: a new workbook with S48 added."""
import csv, shutil
from pathlib import Path
import openpyxl
from openpyxl.styles import Font

SRC = Path("../31_core_gene_verification/Additional_file_1_with_S47.xlsx")
DST = Path("Additional_file_1_with_S48.xlsx")
RES = Path("results")

def tsv(p):
    with open(p, newline="") as fh:
        return [[x.strip() for x in r] for r in csv.reader(fh, delimiter="\t") if r]

def num(v):
    for f in (int, float):
        try:
            return f(v)
        except (ValueError, TypeError):
            pass
    return v

shutil.copy(SRC, DST)
wb = openpyxl.load_workbook(DST)
assert "S47" in wb.sheetnames, "S47 missing from source workbook"
if "S48" in wb.sheetnames:
    del wb["S48"]
ws = wb.create_sheet("S48")
B = Font(bold=True)
row = 1

def block(title, header, rows):
    global row
    ws.cell(row=row, column=1, value=title).font = B
    row += 1
    for j, h in enumerate(header, 1):
        ws.cell(row=row, column=j, value=h).font = B
    row += 1
    for r in rows:
        for j, v in enumerate(r, 1):
            ws.cell(row=row, column=j, value=num(v))
        row += 1
    row += 1

d = tsv(RES / "SUMMARY_DNA.tsv")
block("S48a. DNA-level alignment of each plasmid against every replicon of the other "
      "species (MUMmer4 v4.0.0rc1, nucmer --maxmatch, delta-filter -1). Two filters: "
      "80% identity and 1 kb (as in Methods), and 70% identity and 200 bp (relaxed). "
      "A. marinus Plasmid 1 is the positive control.",
      ["plasmid", "target", "filter", "alignment blocks", "aligned bp (merged)",
       "% of plasmid covered", "length-weighted mean identity (%)"], d[1:])

per = tsv(RES / "per_protein_best_hits.tsv")
summ = {}
for r in per[1:]:
    p, rep, ident = r[0], r[3], r[5]
    s = summ.setdefault(p, {"n": 0, "none": 0, "30-50": 0, "50-70": 0, ">=70": 0,
                            "reps": {}})
    s["n"] += 1
    if rep == "no partner":
        s["none"] += 1
        continue
    x = float(ident)
    b = ">=70" if x >= 70 else "50-70" if x >= 50 else "30-50"
    s[b] += 1
    s["reps"][rep] = s["reps"].get(rep, 0) + 1
rows = []
for p, s in summ.items():
    reps = "; ".join(f"{k}: {v}" for k, v in sorted(s["reps"].items()))
    rows.append([p, s["n"], s["none"], s["30-50"], s["50-70"], s[">=70"], reps])
block("S48b. Protein-level homologs of each plasmid's proteins in the full proteome of "
      "the other species (BLASTP, BLAST+ v2.16.0, E-value 1e-5). A homolog is the best "
      "hit covering at least 50% of the query at 30% identity or higher.",
      ["plasmid", "proteins", "no homolog", "homolog 30-50% identity",
       "homolog 50-70% identity", "homolog >=70% identity", "replicon of best homolog"],
      rows)

block("S48c. Best homolog of every plasmid protein (per-protein detail for S48b).",
      ["plasmid", "locus tag", "product", "replicon of best homolog", "homolog locus tag",
       "identity (%)", "query coverage (%)"], per[1:])

for col, w in zip("ABCDEFG", (26, 30, 22, 16, 18, 18, 34)):
    ws.column_dimensions[col].width = w

rd = wb["README"]
r = rd.max_row + 1
rd.cell(row=r, column=1, value="S48")
rd.cell(row=r, column=2, value="Plasmid comparison against the whole genome of the other species")
rd.cell(row=r, column=3, value=(
    "Each plasmid aligned (MUMmer4 v4.0.0rc1) against every replicon of the other species, "
    "and each plasmid protein searched (BLASTP, BLAST+ v2.16.0) against the full proteome "
    "of the other species. A. marinus Plasmid 1 serves as a positive control."))

wb.save(DST)
print(f"wrote {DST}: {len(wb.sheetnames)} sheets, S48 has {ws.max_row} rows")
