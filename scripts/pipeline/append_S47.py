#!/usr/bin/env python3
"""Append Table S47 to Additional_file_1.xlsx from the verification output.
The input workbook is not edited; a new file is written alongside it."""
import csv, shutil
from pathlib import Path
import openpyxl
from openpyxl.styles import Font

SRC = Path("Additional_file_1.xlsx")
DST = Path("Additional_file_1_with_S47.xlsx")
RES = Path("results_final")
BLA = Path("results/blast")
SP = {"J17J1": "A. marinus", "J31B2": "A. confluentis"}

def tsv(p):
    with open(p, newline="") as fh:
        return [[x.strip() for x in r] for r in csv.reader(fh, delimiter="\t") if r]

def num(v):
    for f in (int, float):
        try:
            return f(v)
        except (ValueError, TypeError):
            pass
    return v

shutil.copy(SRC, DST)
wb = openpyxl.load_workbook(DST)
if "S47" in wb.sheetnames:
    del wb["S47"]
ws = wb.create_sheet("S47")
B = Font(bold=True)
row = 1

def block(title, header, rows):
    global row
    ws.cell(row=row, column=1, value=title).font = B
    row += 1
    for j, h in enumerate(header, 1):
        ws.cell(row=row, column=j, value=h).font = B
    row += 1
    for r in rows:
        for j, v in enumerate(r, 1):
            ws.cell(row=row, column=j, value=num(v))
        row += 1
    row += 1

rows = []
for s in SP:
    rows += [[SP[s]] + r for r in tsv(RES / f"{s}_tRNA_by_replicon.tsv")[1:]]
block("S47a. tRNA genes by replicon. Anticodons from tRNAscan-SE via Bakta v1.12.0; "
      "tRNA-Xxx denotes a predicted tRNA of undetermined isotype.",
      ["species"] + tsv(RES / "J17J1_tRNA_by_replicon.tsv")[0], rows)

rows = []
for s in SP:
    rows += [[SP[s]] + r for r in tsv(RES / f"{s}_aaRS_by_replicon.tsv")[1:]]
block("S47b. Aminoacyl-tRNA synthetases by replicon (Bakta v1.12.0 annotation).",
      ["species"] + tsv(RES / "J17J1_aaRS_by_replicon.tsv")[0], rows)

block("S47c. BLASTP of the A. marinus chromid prolyl-tRNA synthetase against all "
      "chromosomal proteins (BLAST+ v2.16.0, E-value 1e-3). The single hit is the "
      "chromosomal threonyl-tRNA synthetase, a class II paralog, not a second ProS.",
      ["query", "chromosomal subject", "identity (%)", "alignment length",
       "query coverage (%)", "E-value", "bit score"],
      tsv(BLA / "proS_vs_chr.tsv"))

block("S47d. BLASTN of the A. marinus chromid tRNA-Ser(GCU) and elongator tRNA-Met(CAU) "
      "genes against the chromosome sequence (BLAST+ v2.16.0, E-value 1e-3). All hits "
      "are partial matches to other chromosomal tRNA genes; neither query has a "
      "full-length chromosomal copy.",
      ["query", "identity (%)", "alignment length", "query length", "E-value",
       "bit score", "chromosome start", "chromosome end"],
      tsv(BLA / "trna_vs_chr.tsv"))

rows = []
for s in SP:
    d = tsv(RES / f"{s}_codon_counts_by_replicon.tsv")
    hdr = d[0]
    for r in d[1:]:
        if r[0] in ("AGC", "AGT", "ATG"):
            for j in range(1, len(hdr) - 1):
                rows.append([SP[s], r[0], hdr[j], r[j]])
block("S47e. Usage of the AGC, AGT and ATG codons across all coding sequences of each replicon.",
      ["species", "codon", "replicon", "codon count"], rows)

for col, w in zip("ABCDEFGHI", (16, 16, 16, 12, 12, 10, 14, 12, 22)):
    ws.column_dimensions[col].width = w

rd = wb["README"]
r = rd.max_row + 1
rd.cell(row=r, column=1, value="S47")
rd.cell(row=r, column=2, value="Replicon assignment of core decoding functions")
rd.cell(row=r, column=3, value=(
    "tRNA genes by anticodon, aminoacyl-tRNA synthetases, and codon usage assigned to "
    "their replicon of origin from the Bakta v1.12.0 annotation, with BLAST+ v2.16.0 "
    "confirmation that the three A. marinus chromid-only decoding functions have no "
    "unannotated chromosomal copy."))

wb.save(DST)
print(f"wrote {DST}: {len(wb.sheetnames)} sheets, S47 has {ws.max_row} rows")
